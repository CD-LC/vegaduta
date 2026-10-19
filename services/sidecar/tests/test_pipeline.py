"""
First-ever sidecar tests (RAG audit, 2026-07-29). Scope is deliberately the
pure-logic surfaces this audit touched or depends on - chunk math, the
embedding model registry's invariants, and the extraction fixes (DOCX
tables, scanned-PDF fallback error) - not the model-loading/OCR paths, which
need multi-hundred-MB downloads or the tesseract system binary and belong in
an integration suite, not unit tests.

Run from sidecar/:  pytest tests/
(pytest itself is a dev-only dependency - see requirements-dev.txt; it is
deliberately NOT in requirements.txt so it never ships in the image.)
"""
import io

import pytest

from app.chunking import chunk_text
from app.embeddings import (
    DEFAULT_MODEL_KEY,
    MODEL_REGISTRY,
    STORAGE_DIMENSIONS,
    resolve_model_key,
)
from app.extraction import (
    _MIN_PAGE_TEXT_CHARS,
    _extension_of,
    _extract_docx,
    _extract_html,
    _extract_pdf,
)

# --- chunking ---------------------------------------------------------------

class TestChunkText:
    def test_short_text_is_a_single_chunk(self):
        assert chunk_text("one two three", chunk_size=150, chunk_overlap=30) == ["one two three"]

    def test_empty_and_whitespace_only_text_yield_no_chunks(self):
        assert chunk_text("") == []
        assert chunk_text("   \n\t  ") == []

    def test_overlap_must_be_smaller_than_chunk_size(self):
        with pytest.raises(ValueError):
            chunk_text("some words here", chunk_size=10, chunk_overlap=10)

    def test_overlap_is_exact_word_for_word(self):
        # The module docstring promises: "chunk_overlap words at the tail of
        # chunk N reappear at the head of chunk N+1, exactly, every time."
        words = [f"w{i}" for i in range(10)]
        chunks = chunk_text(" ".join(words), chunk_size=4, chunk_overlap=2)

        assert chunks == [
            "w0 w1 w2 w3",
            "w2 w3 w4 w5",
            "w4 w5 w6 w7",
            "w6 w7 w8 w9",
        ]
        for previous, current in zip(chunks, chunks[1:], strict=False):  # noqa: RUF007
            assert previous.split()[-2:] == current.split()[:2]

    def test_every_word_appears_and_none_are_split(self):
        words = [f"word{i}" for i in range(37)]
        chunks = chunk_text(" ".join(words), chunk_size=10, chunk_overlap=3)

        seen = set()
        for chunk in chunks:
            seen.update(chunk.split())
        assert seen == set(words)


# --- embeddings registry (no model loading) ---------------------------------

class TestModelRegistry:
    def test_default_key_is_registered_and_unchanged(self):
        # Changing the default would silently invalidate existing
        # collections' stored vectors - locked down on purpose.
        assert DEFAULT_MODEL_KEY == "minilm-l6-v2"
        assert DEFAULT_MODEL_KEY in MODEL_REGISTRY

    def test_multilingual_model_is_registered_with_padding_safe_dimensions(self):
        # RAG audit fix: the Indic-capable model must exist in BOTH this
        # registry and EmbeddingModelCatalog.java (they mirror each other).
        repo_id, dims = MODEL_REGISTRY["multilingual-minilm"]
        assert repo_id == "sentence-transformers/paraphrase-multilingual-MiniLM-L12-v2"
        assert dims == 384

    def test_no_model_exceeds_storage_dimensions(self):
        # Zero-padding up to STORAGE_DIMENSIONS only works when every native
        # width fits - a wider model would corrupt the shared vector(768)
        # column's similarity math (embeddings.py raises at runtime, but the
        # registry should never contain such an entry in the first place).
        for key, (_repo, dims) in MODEL_REGISTRY.items():
            assert dims <= STORAGE_DIMENSIONS, f"{key} native dims {dims} exceed {STORAGE_DIMENSIONS}"  # noqa: E501

    def test_resolve_model_key_defaults_and_validates(self):
        assert resolve_model_key(None) == DEFAULT_MODEL_KEY
        assert resolve_model_key("") == DEFAULT_MODEL_KEY
        assert resolve_model_key("multilingual-minilm") == "multilingual-minilm"
        with pytest.raises(ValueError):
            resolve_model_key("no-such-model")


# --- extraction -------------------------------------------------------------

class TestExtraction:
    def test_extension_of_rejects_unsupported(self):
        with pytest.raises(ValueError):
            _extension_of("malware.exe")

    def test_docx_extracts_paragraphs_and_tables(self):
        # RAG audit fix: table content used to be dropped entirely.
        from docx import Document as DocxDocument

        document = DocxDocument()
        document.add_paragraph("Intro paragraph text")
        table = document.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "Plan"
        table.rows[0].cells[1].text = "Price"
        table.rows[1].cells[0].text = "Premium"
        table.rows[1].cells[1].text = "Rs 999"
        buffer = io.BytesIO()
        document.save(buffer)

        text = _extract_docx(buffer.getvalue())

        assert "Intro paragraph text" in text
        assert "Plan\tPrice" in text
        assert "Premium\tRs 999" in text

    def test_docx_skips_fully_empty_table_rows(self):
        from docx import Document as DocxDocument

        document = DocxDocument()
        table = document.add_table(rows=2, cols=2)
        table.rows[0].cells[0].text = "only row with content"
        buffer = io.BytesIO()
        document.save(buffer)

        text = _extract_docx(buffer.getvalue())

        lines = [line for line in text.split("\n") if line.strip()]
        assert lines == ["only row with content\t"] or lines == ["only row with content"]

    def test_corrupt_docx_raises_a_clear_value_error(self):
        with pytest.raises(ValueError, match="Could not read DOCX"):
            _extract_docx(b"this is not a zip archive")

    def test_html_strips_scripts_and_styles(self):
        html = b"<html><head><style>body{color:red}</style></head><body><script>alert(1)</script><p>Visible content</p></body></html>"  # noqa: E501

        text = _extract_html(html)

        assert "Visible content" in text
        assert "alert(1)" not in text
        assert "color:red" not in text

    def test_textless_pdf_without_images_raises_actionable_error(self):
        # RAG audit fix: an image-free, text-layer-free PDF used to bubble up
        # as a generic "no readable text" 400 - it must now say what the
        # actual problem (and remedy) is. A blank page has no text layer and
        # no embedded images, so the OCR fallback finds nothing either.
        from pypdf import PdfWriter

        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        buffer = io.BytesIO()
        writer.write(buffer)

        with pytest.raises(ValueError, match="no text layer"):
            _extract_pdf(buffer.getvalue())

    def test_corrupt_pdf_raises_a_clear_value_error(self):
        with pytest.raises(ValueError, match="Could not read PDF"):
            _extract_pdf(b"%PDF-not really a pdf at all")

    def test_min_page_text_threshold_is_small(self):
        # The OCR fallback must only kick in for effectively-textless pages;
        # a normal paragraph of text must always clear the threshold.
        assert 0 < _MIN_PAGE_TEXT_CHARS <= 50
