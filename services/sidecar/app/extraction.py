"""
Knowledge Stage F3a/F3b: text extraction for non-Markdown upload formats.

Each extractor takes raw file bytes and returns plain text - downstream,
this feeds the exact same chunk_text()/embed_texts() pipeline plain
Markdown already went through (main.py's /extract endpoint is the only new
thing; everything past it is unchanged), so nothing else needs to know
which format a tenant originally uploaded.

Deliberately conservative on failure: a corrupt/malformed file raises
ValueError with a specific message rather than letting a raw library
traceback escape - main.py's /extract endpoint turns a ValueError into a
400, not a 500, since a bad upload is a client problem, not a sidecar bug.
This also matters for the Stage F Build Plan's stated risk posture: file
parsers are a real, historically-exploited attack surface (OOXML zip
bombs, PDF parser CVEs) - extraction running here, in the already-separate
sidecar process, means a parser crash/hang never takes down the core API
serving every other tenant's requests; it only takes down this one
request.
"""
import base64
import io
import logging
import os

import pytesseract
import requests
from bs4 import BeautifulSoup
from docx import Document
from PIL import Image
from pypdf import PdfReader

logger = logging.getLogger(__name__)

# Local (non-Docker) dev on Windows: installing tesseract-ocr doesn't put it
# on PATH the way `apt-get install`/`brew install` do on Linux/macOS, and
# there's no one standard install location. Rather than force everyone onto
# PATH edits, honor an explicit override if set - see README.md's sidecar
# setup section ("1b. Install Tesseract OCR") for how a person actually sets
# this. No-op (falls through to pytesseract's own PATH lookup, which is what
# every other platform already relies on) when unset.
_tesseract_cmd_override = os.environ.get("TESSERACT_CMD")
if _tesseract_cmd_override:
    pytesseract.pytesseract.tesseract_cmd = _tesseract_cmd_override

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".html", ".htm", ".png", ".jpg", ".jpeg", ".webp"}

# Stage F3b - OCR only, not image captioning/vision-embedding. Covers the
# highest-value real case (screenshots, scanned documents, receipts,
# photographed whiteboards) cheaply and locally, consistent with this
# codebase's "no hosted API required" embeddings philosophy - see the
# Stage F Build Plan's Section 5 for why full image captioning is
# deliberately deferred, not silently dropped.
_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}

# OCR provider abstraction (tenant-configurable cloud OCR, "give an option
# to use any OCR AI service like google ocr...else use open source OCR
# reading library"). Only GOOGLE_VISION has a real call implemented here -
# the other two enum values OcrProvider models on the Java side
# (AWS_TEXTRACT, AZURE_DOCUMENT_INTELLIGENCE) fall through to Tesseract with
# a clear log line below, the same "never silently do nothing" posture as
# every other not-yet-wired integration point in this codebase.
_GOOGLE_VISION_ANNOTATE_URL = "https://vision.googleapis.com/v1/images:annotate"


def extract_text(filename: str, content: bytes, ocr_provider: str | None = None, ocr_api_key: str | None = None) -> str:  # noqa: E501
    extension = _extension_of(filename)
    if extension == ".pdf":
        return _extract_pdf(content)
    if extension == ".docx":
        return _extract_docx(content)
    if extension in _IMAGE_EXTENSIONS:
        return _extract_image(content, ocr_provider, ocr_api_key)
    # .html and .htm share one extractor
    return _extract_html(content)


def _extension_of(filename: str) -> str:
    lowered = filename.lower()
    for ext in SUPPORTED_EXTENSIONS:
        if lowered.endswith(ext):
            return ext
    raise ValueError(f"Unsupported file extension for extraction: {filename}")


# RAG audit fix (2026-07-29): a PDF page whose text layer yields fewer than
# this many non-whitespace characters is treated as effectively textless
# (scanned/image-only) and gets the OCR fallback below. Deliberately a small
# threshold - real text-layer pages practically always clear it, while
# scanner apps sometimes stamp a few stray characters (page numbers, a date)
# onto otherwise image-only pages.
_MIN_PAGE_TEXT_CHARS = 25


def _extract_pdf(content: bytes) -> str:
    """
    Text-layer extraction via pypdf, with a per-page OCR fallback for
    scanned/image-only pages (RAG audit fix, 2026-07-29 - previously an
    image-only PDF silently yielded blank text and a misleading 400).

    The fallback deliberately uses only dependencies already in
    requirements.txt: pypdf's `page.images` pulls each page's EMBEDDED image
    streams (a scanned page is, almost by definition, one full-page embedded
    image), Pillow decodes them, and the same Tesseract wiring
    _extract_image already relies on OCRs them. This is NOT page
    rasterization - a page whose "text" is pure vector drawing (no text
    layer, no embedded raster images) still yields nothing, because true
    rasterization needs a renderer dependency (pdf2image+poppler or
    pymupdf) this sidecar intentionally doesn't ship yet; that case now
    raises a specific, actionable error instead of the generic "no readable
    text" one.
    """
    try:
        reader = PdfReader(io.BytesIO(content))
    except Exception as e:
        raise ValueError(f"Could not read PDF: {e}") from e

    pages: list[str] = []
    for page in reader.pages:
        text = page.extract_text() or ""
        if len(text.strip()) < _MIN_PAGE_TEXT_CHARS:
            ocr_text = _ocr_pdf_page_images(page)
            # Keep whichever source said more - a nearly-blank text layer
            # plus a full-page scan should surface the scan's OCR text, but
            # OCR noise must never REPLACE a real (if short) text layer
            # with less than it had.
            if len(ocr_text.strip()) > len(text.strip()):
                text = ocr_text
        pages.append(text)

    joined = "\n\n".join(pages)
    if not joined.strip():
        raise ValueError(
            "This PDF has no text layer and no OCR-readable embedded page images. "
            "If it is a scanned document, re-export it with OCR applied, or upload the "
            "page images directly (.png/.jpg) - image uploads are OCR'd here."
        )
    return joined


def _ocr_pdf_page_images(page) -> str:
    """
    OCRs every embedded image on a pypdf page object, returning the
    concatenated text (empty string when there are no images, none decode,
    or the Tesseract binary is missing - per-image failures are logged and
    skipped rather than failing the whole document, since one exotic image
    filter pypdf can't decode, e.g. JBIG2, shouldn't sink the pages it CAN
    read).
    """
    try:
        images = page.images
    except Exception as e:
        logger.debug("Could not enumerate embedded images on PDF page: %s", e)
        return ""

    texts: list[str] = []
    for image_file in images:
        try:
            image = Image.open(io.BytesIO(image_file.data))
            image.load()
        except Exception as e:
            logger.debug("Could not decode embedded PDF image %s: %s", getattr(image_file, "name", "?"), e)  # noqa: E501
            continue
        try:
            texts.append(pytesseract.image_to_string(image))
        except pytesseract.TesseractNotFoundError:
            # Deployment problem, not a document problem - same distinction
            # _extract_image draws. For the PDF fallback path we log rather
            # than raise: a text-layer PDF must never start failing just
            # because OCR (a best-effort enhancement here) is unavailable.
            logger.warning("tesseract-ocr binary not found - scanned-PDF OCR fallback is unavailable on this deployment")  # noqa: E501
            return ""
        except Exception as e:
            logger.debug("OCR failed for an embedded PDF image: %s", e)
    return "\n".join(t for t in texts if t.strip())


def _extract_docx(content: bytes) -> str:
    try:
        document = Document(io.BytesIO(content))
    except Exception as e:
        raise ValueError(f"Could not read DOCX: {e}") from e
    parts = [paragraph.text for paragraph in document.paragraphs]
    # RAG audit fix (2026-07-29): table content was dropped entirely
    # (document.paragraphs never descends into tables), so a DOCX whose
    # substance lives in tables - price lists, spec sheets, policy matrices -
    # ingested as near-empty. python-docx exposes tables directly; one
    # tab-joined line per row keeps cell adjacency visible to chunking
    # without inventing a layout format.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                parts.append("\t".join(cells))
    return "\n".join(parts)


def _extract_image(content: bytes, ocr_provider: str | None = None, ocr_api_key: str | None = None) -> str:  # noqa: E501
    try:
        image = Image.open(io.BytesIO(content))
        # Force a decode now, inside this try - Pillow is lazy by default
        # (Image.open only reads the header), so a truncated/corrupt image
        # body would otherwise surface as an opaque error later, from
        # pytesseract, instead of a clear "could not read image" here.
        image.load()
    except Exception as e:
        raise ValueError(f"Could not read image: {e}") from e

    if ocr_provider == "GOOGLE_VISION" and ocr_api_key:
        try:
            return _call_google_vision(content, ocr_api_key)
        except Exception as e:
            # A cloud-provider failure (bad key, quota, network) is not the
            # uploaded image's fault, and per OcrProviderConfigService's
            # javadoc a bad key is expected to surface here, at ingestion
            # time, not earlier - fall back to local Tesseract rather than
            # failing the whole upload, but log loudly so it's visible in
            # the sidecar's own logs which tenant/provider combo is broken.
            logger.warning("Google Vision OCR failed, falling back to Tesseract: %s", e)
    elif ocr_provider in ("AWS_TEXTRACT", "AZURE_DOCUMENT_INTELLIGENCE"):
        # Modeled on the Java side (OcrProvider enum) but no sidecar call is
        # implemented yet for these two - explicit log line, not a silent
        # no-op, then fall through to Tesseract exactly as if unconfigured.
        logger.info("OCR provider %s has no sidecar implementation yet; using Tesseract", ocr_provider)  # noqa: E501

    try:
        text = pytesseract.image_to_string(image)
    except pytesseract.TesseractNotFoundError as e:
        # Distinct from every other extractor's failure mode: this means
        # the sidecar's own environment is missing the tesseract-ocr system
        # binary (see the Dockerfile's comment) - a deployment/ops problem,
        # not something the uploaded file did wrong. Still a ValueError (400)
        # rather than a 500 for consistency with every other extractor here,
        # but the message says so explicitly rather than implying the image
        # itself was bad.
        raise ValueError(
            "OCR is not available on this sidecar deployment (tesseract-ocr binary not found)"
        ) from e
    except Exception as e:
        raise ValueError(f"OCR failed: {e}") from e
    return text


def _call_google_vision(content: bytes, api_key: str) -> str:
    """
    Calls Google Cloud Vision's TEXT_DETECTION feature via its plain REST
    API (not the google-cloud-vision SDK) since auth here is just a
    tenant-supplied API key as a query param - see OcrProvider's javadoc for
    why Google Vision is the only provider with a real call in this pass.
    Raises on any failure; the caller (_extract_image) decides whether that
    means "fall back to Tesseract."
    """
    payload = {
        "requests": [
            {
                "image": {"content": base64.b64encode(content).decode("ascii")},
                "features": [{"type": "TEXT_DETECTION"}],
            }
        ]
    }
    response = requests.post(
        _GOOGLE_VISION_ANNOTATE_URL,
        params={"key": api_key},
        json=payload,
        timeout=30,
    )
    response.raise_for_status()
    body = response.json()
    responses = body.get("responses") or []
    if not responses:
        raise ValueError("Google Vision returned no responses")
    result = responses[0]
    if "error" in result:
        raise ValueError(f"Google Vision error: {result['error'].get('message', 'unknown')}")
    full_text_annotation = result.get("fullTextAnnotation")
    if full_text_annotation and full_text_annotation.get("text"):
        return full_text_annotation["text"]
    # Fall back to the first individual text-annotation entry (the "whole
    # image" summary) if fullTextAnnotation wasn't populated - can happen on
    # images with sparse/no detected text.
    text_annotations = result.get("textAnnotations") or []
    if text_annotations:
        return text_annotations[0].get("description", "")
    return ""


def _extract_html(content: bytes) -> str:
    # "html.parser" is Python's stdlib parser - deliberately not lxml, to
    # avoid adding a C-extension system dependency for a format that
    # doesn't need one.
    try:
        soup = BeautifulSoup(content, "html.parser")
    except Exception as e:
        raise ValueError(f"Could not parse HTML: {e}") from e
    # Script/style content is never useful as "knowledge" text and is
    # commonly huge relative to a page's actual visible content.
    for tag in soup(["script", "style"]):
        tag.decompose()
    return soup.get_text(separator="\n")
